"""'BOM Comparison' sheet of the Cleanup Report.

Standard BOM (first level only) vs. the first-level BOM of the 3D CAD assembly
(top-level occurrences, Occurrence Properties). Nested subassembly content is never
counted. Same strategy as the manual excel_pdf_compare.py: count every item on both
sides, then compare the counts per item.

    Missing in 3D CAD  : in Standard BOM, not in 3D CAD  -> add the part to the CAD
    Extra in 3D CAD    : in 3D CAD, not in Standard BOM  -> hide the part
    Quantity Mismatch  : in both, different quantity
    Match              : in both, same quantity
"""

from openpyxl.styles import Font, PatternFill

from ..solidedge.model import part_number
from .excel import write_table

MISSING = "Missing in 3D CAD"
EXTRA = "Extra in 3D CAD"
MISMATCH = "Quantity Mismatch"
MATCH = "Match"
_ORDER = {MISSING: 0, EXTRA: 1, MISMATCH: 2, MATCH: 3}
_FILLS = {MISSING: "FFC7CE", EXTRA: "FFEB9C", MISMATCH: "FFEB9C", MATCH: "C6EFCE"}
_ACTIONS = {MISSING: "Add part to the 3D / 2D CAD file", EXTRA: "Hide part in 3D CAD",
            MISMATCH: "Check quantity in 3D CAD", MATCH: "-"}

HEADERS = ["Item No", "Status", "Action", "Name (Standard BOM)", "Title (3D CAD)",
           "Expected Quantity (Standard BOM)", "Actual Quantity (3D CAD 1st level)",
           "Total Quantity Missing", "BOM Position", "BOM Revision", "3D CAD Occurrences", "Remarks"]


def _qty(text) -> int:
    try:
        return int(float(str(text).strip() or 0))
    except ValueError:
        return 0


def compare_bom(bom_lines, nodes) -> list[list]:
    """Rows for the comparison table. Only level-1 nodes are used."""
    expected, bom_info = {}, {}
    for b in bom_lines or []:
        key = b.key
        if not key:
            continue
        expected[key] = expected.get(key, 0) + (_qty(b.quantity) if str(b.quantity).strip() else 1)
        info = bom_info.setdefault(key, {"name": b.name, "pos": [], "rev": b.revision, "id": str(b.item_id).strip()})
        if b.position:
            info["pos"].append(str(b.position))

    actual, cad_info = {}, {}
    for n in nodes:
        if n.level != 1:                              # nested subassembly content is excluded
            continue
        key = part_number(n.file or n.name)
        doc_key = part_number(n.doc_number)
        if key not in expected and doc_key and doc_key in expected:
            key = doc_key                             # document number matches the BOM item
        info = cad_info.setdefault(key, {"title": n.title, "occ": [], "rpl_no": 0, "missing": 0})
        info["occ"].append(n.name)
        info["rpl_no"] += n.include_in_bom is False
        info["missing"] += bool(n.file_missing)
        actual[key] = actual.get(key, 0) + 1          # every top-level occurrence counts once

    rows = []
    for key in set(expected) | set(actual):
        exp, act = expected.get(key, 0), actual.get(key, 0)
        if act == 0:
            status = MISSING
        elif exp == 0:
            status = EXTRA
        elif exp == act:
            status = MATCH
        else:
            status = MISMATCH
        b, c = bom_info.get(key, {}), cad_info.get(key, {})
        remarks = []
        if c.get("rpl_no"):
            remarks.append(f"{c['rpl_no']} occurrence(s) with Reports/Parts List = No")
        if c.get("missing"):
            remarks.append(f"{c['missing']} occurrence(s) with missing file")
        rows.append([b.get("id") or key, status, _ACTIONS[status], b.get("name", ""), c.get("title", ""),
                     exp, act, exp - act, ", ".join(b.get("pos", [])), b.get("rev", ""),
                     "; ".join(c.get("occ", [])[:20]) + (" ..." if len(c.get("occ", [])) > 20 else ""),
                     "; ".join(remarks)])
    rows.sort(key=lambda r: (_ORDER[r[1]], str(r[0])))
    return rows


def write_bom_comparison(wb, bom_lines, nodes, bom_name: str = ""):
    ws = wb.create_sheet("BOM Comparison")
    if not bom_lines:
        ws.append(["Standard BOM not available - comparison not possible."])
        ws.cell(row=1, column=1).font = Font(bold=True)
        return
    rows = compare_bom(bom_lines, nodes)
    write_table(ws, HEADERS, rows, widths={"Name (Standard BOM)": 40, "Title (3D CAD)": 40,
                                          "3D CAD Occurrences": 45, "Remarks": 50, "Action": 32})
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=2).value
        if status in _FILLS:
            fill = PatternFill("solid", fgColor=_FILLS[status])
            ws.cell(row=r, column=2).fill = fill
            ws.cell(row=r, column=8).fill = fill

    # summary + legend to the right of the table (like the manual comparison report)
    col = len(HEADERS) + 2
    count = {s: sum(1 for x in rows if x[1] == s) for s in _ORDER}
    lines = [
        ("STANDARD BOM vs 3D CAD (1st level)", None),
        (f"Standard BOM: {bom_name}", None),
        ("Only top-level occurrences are compared; parts inside subassemblies are excluded.", None),
        ("", None),
        (f"Parts Missing in 3D CAD: {count[MISSING]}  -> add to 3D / 2D CAD file", MISSING),
        (f"Extra Parts in 3D CAD: {count[EXTRA]}  -> hide in 3D CAD", EXTRA),
        (f"Quantity Mismatch: {count[MISMATCH]}", MISMATCH),
        (f"Match: {count[MATCH]}", MATCH),
        ("", None),
        ("--- LEGEND: TOTAL QUANTITY MISSING ---", None),
        ("Positive (+) : Missing in 3D CAD (Standard BOM has more)", None),
        ("Negative (-) : Extra in 3D CAD (3D CAD has more)", None),
    ]
    for i, (text, status) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=col, value=text)
        if i == 1 or text.startswith("---"):
            cell.font = Font(bold=True)
        if status:
            cell.fill = PatternFill("solid", fgColor=_FILLS[status])
    ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 70
