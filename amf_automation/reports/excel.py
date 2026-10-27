"""Shared openpyxl formatting."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILLS = {
    "PASS": "C6EFCE", "SUCCESS": "C6EFCE", "DOWNLOADED": "C6EFCE", "YES": "C6EFCE", "KEEP": "C6EFCE",
    "FAIL": "FFC7CE", "FAILED": "FFC7CE", "NO": "FFC7CE", "ERROR": "FFC7CE",
    "WARN": "FFEB9C", "REVIEW": "FFEB9C", "PARTIAL": "FFEB9C", "NOT AVAILABLE": "FFEB9C",
    "SKIPPED": "EDEDED", "N/A": "EDEDED",
}


def status_fill(value):
    if not isinstance(value, str):
        return None
    v = value.strip().upper()
    for key, color in STATUS_FILLS.items():
        if v == key or v.startswith(key + " ") or v.startswith(key + "("):
            return PatternFill("solid", fgColor=color)
    return None


def write_table(ws, headers: list[str], rows: list[list], status_cols: set[str] | None = None,
                widths: dict[str, int] | None = None, max_width: int = 60):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
    status_idx = {i for i, h in enumerate(headers) if status_cols and h in status_cols}
    for row in rows:
        ws.append([_cell(v) for v in row])
        r = ws.max_row
        for i in range(len(headers)):
            cell = ws.cell(row=r, column=i + 1)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if i in status_idx:
                fill = status_fill(cell.value)
                if fill:
                    cell.fill = fill
    for i, h in enumerate(headers, start=1):
        letter = get_column_letter(i)
        if widths and h in widths:
            ws.column_dimensions[letter].width = widths[h]
            continue
        longest = max([len(str(h))] + [len(str(ws.cell(row=r, column=i).value or ""))
                                       for r in range(2, min(ws.max_row, 300) + 1)])
        ws.column_dimensions[letter].width = min(max(10, longest + 2), max_width)
    ws.freeze_panes = "B2" if len(headers) > 3 else "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def write_key_values(ws, pairs: list[tuple[str, object]], title: str | None = None):
    if title:
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14)
        ws.append([])
    for k, v in pairs:
        ws.append([k, _cell(v)])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        fill = status_fill(ws.cell(row=ws.max_row, column=2).value)
        if fill:
            ws.cell(row=ws.max_row, column=2).fill = fill
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 90


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (int, float, str)):
        return v
    return str(v)
