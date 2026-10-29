"""Summary report for a whole cycle (all serial numbers) + permanent history workbook."""

from pathlib import Path

import openpyxl

from ..models import PROCESS1_STEPS, PROCESS2_STEPS, SerialRecord
from ..utils import iso
from .excel import write_key_values, write_table

# (column header, record.data key). Contains every column of the original summary_report.xlsx
# (Row Status -> "Status in Bluestar", Revision -> "Revision in Bluestar", PDF/BOM Status ->
# "2D Draft Status"/"Standard BOM Status", Timestamp -> "Timestamp of Completion") plus all
# step / checkpoint columns of the standard process.
SUMMARY_COLUMNS = [
    ("SR_NO", "serial"),
    ("Overall Status", "overall_status"),
    ("Input Row", "row"),
    ("Input Validation", "validation"),
    ("VPN Status", "vpn_status"),
    ("Object ID", "object_id"),
    ("Object Name", "object_name"),
    ("Revision Filter Used", "revision_filter"),
    ("Status in Bluestar", "bluestar_status"),
    ("Revision in Bluestar", "revision"),
    ("Download Status", "download_status"),
    ("CAD File", "cad_file"),
    ("CAD Status", "cad_status"),
    ("CAD Files Downloaded", "files_downloaded"),
    ("CAD Download Size", "download_size"),
    ("PDF Path", "pdf_path"),
    ("2D Draft Status", "pdf_status"),
    ("BOM Path", "bom_path"),
    ("Standard BOM Status", "bom_status"),
    ("Standard BOM Lines", "bom_lines"),
    ("Working Folder", "working_folder"),
    ("Move to Working Folder Status", "working_status"),
    ("Process 1 Status", "p1_status"),
    ("Top-level Assembly", "top_assembly"),
    ("Components Analyzed", "n_analyzed"),
    ("Components Retained", "n_retained"),
    ("Components Hidden", "n_hidden"),
    ("Hidden - Not in Standard BOM", "n_hidden_bom"),
    ("Hidden - Missing Occurrence/Part", "n_hidden_missing"),
    ("Electrical Identified", "n_electrical"),
    ("Electrical Retained", "n_electrical_kept"),
    ("Electrical Unhidden (rule override)", "n_electrical_unhidden"),
    ("Motors Identified", "n_motors"),
    ("Motors Coloured Red", "n_motors_red"),
    ("Motors Not Coloured", "n_motors_failed"),
    ("Protected Components Kept", "n_protected_kept"),
    ("Assemblies Hidden (Hide All)", "n_hide_all"),
    ("Laser Sensor Parts", "n_laser_files"),
    ("Laser Beams Removed", "n_laser_beams_removed"),
    ("Laser Sensors Need Review", "n_laser_review"),
    ("Laser Beam Removal Status", "laser_status"),
    ("Ambiguous - Need Review", "n_review"),
    ("Missing Files Found", "missing_found"),
    ("Missing Files Repaired", "missing_repaired"),
    ("Missing Files Not in Bluestar", "missing_not_in_bluestar"),
    ("Z-up Orientation", "orientation"),
    ("Checkpoint-1 Status (Show All)", "cp1"),
    ("Checkpoint-2 Status (Missing Files)", "cp2"),
    ("Checkpoint-3 Status (Z-up)", "cp3"),
    ("Checkpoint-4 Status (Part Style)", "cp4"),
    ("Checkpoint-5 Status (BOM / Hide)", "cp5"),
    ("Checkpoint-6 Status (Clean up)", "cp6"),
    ("Checkpoint-7 Status (STEP Export)", "cp7"),
    ("STEP Available", "step_available"),
    ("STEP Path", "step_path"),
    ("Cleanup Report", "cleanup_report"),
    ("Process 2 Status", "p2_status"),
    ("Manual Action Required", "manual_action"),
    ("Error", "error"),
    ("Start Time", "start_time"),
    ("Timestamp of Completion", "timestamp"),
    ("Duration (min)", "duration_min"),
]

STATUS_HEADERS = {"Download Status", "CAD Status", "Overall Status", "VPN Status",
                  "Status in Bluestar", "2D Draft Status", "Standard BOM Status",
                  "Move to Working Folder Status", "Process 1 Status", "Process 2 Status", "Z-up Orientation",
                  "STEP Available", "Input Validation", "Laser Beam Removal Status"} | {h for h, _ in SUMMARY_COLUMNS if h.startswith("Checkpoint")}


def record_row(rec: SerialRecord) -> list:
    d = dict(rec.data)
    d.setdefault("serial", rec.serial)
    d.setdefault("row", rec.row)
    d.setdefault("validation", rec.validation)
    if rec.manual_actions and not d.get("manual_action"):
        d["manual_action"] = " | ".join(rec.manual_actions[:15]) + (" ..." if len(rec.manual_actions) > 15 else "")
    return [d.get(key, "") for _h, key in SUMMARY_COLUMNS]


def write_summary(path: Path, run, records: list[SerialRecord]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    write_table(ws, [h for h, _ in SUMMARY_COLUMNS], [record_row(r) for r in records], STATUS_HEADERS)

    ws2 = wb.create_sheet("Step Flags")
    all_steps = PROCESS1_STEPS + PROCESS2_STEPS
    rows = []
    for r in records:
        for key, name in all_steps:
            f = r.steps.get(key)
            if f is None:
                continue
            rows.append([r.serial, key, name, f.status, f.message, iso(f.started), iso(f.finished)])
    write_table(ws2, ["SR_NO", "Step", "Description", "Status", "Message", "Started", "Finished"], rows,
                {"Status"}, widths={"Message": 90, "Description": 55})

    ws3 = wb.create_sheet("Step Matrix")
    headers = ["SR_NO", "Overall Status"] + [k for k, _ in all_steps]
    matrix = [[r.serial, r.data.get("overall_status", "")] +
              [(r.steps[k].status if k in r.steps else "") for k, _ in all_steps] for r in records]
    write_table(ws3, headers, matrix, set(headers[1:]))

    ws4 = wb.create_sheet("Cycle Info")
    total = len(records)
    count = lambda pred: sum(1 for r in records if pred(r))  # noqa: E731
    write_key_values(ws4, [
        ("Cycle ID", run.cycle_id),
        ("Input file", str(run.input_file)),
        ("Started", iso(run.started)),
        ("Finished", iso(run.finished) or "running"),
        ("Serial numbers in input", total),
        ("Valid (distinct, new)", count(lambda r: r.is_valid)),
        ("Duplicates", count(lambda r: r.validation.startswith("Duplicate"))),
        ("Invalid", count(lambda r: r.validation.startswith("Invalid"))),
        ("Already processed earlier", count(lambda r: r.validation.startswith("Already"))),
        ("SUCCESS", count(lambda r: r.data.get("overall_status") == "SUCCESS")),
        ("SUCCESS - review needed", count(lambda r: r.data.get("overall_status") == "SUCCESS - REVIEW")),
        ("FAILED", count(lambda r: r.data.get("overall_status") == "FAILED")),
        ("Not processed", count(lambda r: r.data.get("overall_status") in ("NOT PROCESSED", "", None) and r.is_valid)),
        ("Cycle aborted", run.aborted_reason or "No"),
        ("Step legend", "PASS = done | WARN = done, needs review | FAIL = stopped | SKIPPED = not run"),
    ], title="AMF Solid Edge Clean Up Automation - Cycle Summary")
    _save(wb, path)


def write_history(path: Path, history: list[dict]):
    """One workbook with every serial ever processed (all cycles)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "History"
    headers = ["Cycle ID"] + [h for h, _ in SUMMARY_COLUMNS]
    rows = [[h.get("cycle_id", "")] + [h.get(key, "") for _hh, key in SUMMARY_COLUMNS] for h in history]
    write_table(ws, headers, rows, STATUS_HEADERS)
    _save(wb, path)


def _save(wb, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:      # file open in Excel -> save a copy instead of crashing
        alt = path.with_name(path.stem + "_new" + path.suffix)
        wb.save(alt)
